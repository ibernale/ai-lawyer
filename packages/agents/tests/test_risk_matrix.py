"""Unit tests for RiskMatrixGenerator and build_risk_matrix_xlsx."""

from __future__ import annotations

import json
from datetime import UTC, datetime
from unittest.mock import MagicMock

import pytest
from anthropic.types import TextBlock
from lex_agents_agents.risk_matrix import (
    RiskItem,
    RiskLevel,
    RiskMatrix,
    RiskMatrixGenerator,
    _derive_risk_level,
    _extract_json_array,
)
from lex_agents_agents.risk_matrix_xlsx import build_risk_matrix_xlsx

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_consult_response(answer: str = "Análisis de prueba.") -> MagicMock:
    resp = MagicMock()
    resp.trace_id = "test-trace-001"
    resp.answer = answer
    resp.citations = []
    return resp


def _make_client(llm_text: str) -> MagicMock:
    """Return a mock AnthropicClientWrapper whose messages_create returns *llm_text*."""
    text_block = MagicMock(spec=TextBlock, text=llm_text)
    api_response = MagicMock()
    api_response.content = [text_block]
    client = MagicMock()
    client.messages_create.return_value = api_response
    return client


_SAMPLE_JSON = json.dumps([
    {
        "regulation": "CRR Art. 92",
        "obligation": "Mantener ratio de capital mínimo del 8%",
        "likelihood": "Alta",
        "impact": "Alto",
        "mitigation": "Monitoreo diario del ratio CET1",
    },
    {
        "regulation": "DORA Art. 5",
        "obligation": "Gestión del riesgo TIC",
        "likelihood": "Media",
        "impact": "Medio",
        "mitigation": "Implementar marco DORA antes de enero 2025",
    },
    {
        "regulation": "AMLD6 Art. 18",
        "obligation": "Due diligence reforzada para PEPs",
        "likelihood": "Baja",
        "impact": "Bajo",
        "mitigation": "Revisar procedimientos de onboarding",
    },
])


# ---------------------------------------------------------------------------
# _derive_risk_level
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "likelihood,impact,expected",
    [
        ("Alta", "Alto", RiskLevel.ALTO),
        ("Alta", "Medio", RiskLevel.ALTO),
        ("Alta", "Bajo", RiskLevel.MEDIO),
        ("Media", "Alto", RiskLevel.ALTO),
        ("Media", "Medio", RiskLevel.MEDIO),
        ("Media", "Bajo", RiskLevel.BAJO),
        ("Baja", "Alto", RiskLevel.MEDIO),
        ("Baja", "Medio", RiskLevel.BAJO),
        ("Baja", "Bajo", RiskLevel.BAJO),
    ],
)
def test_derive_risk_level(likelihood: str, impact: str, expected: RiskLevel) -> None:
    assert _derive_risk_level(likelihood, impact) == expected


def test_derive_risk_level_unknown_key_defaults_medio() -> None:
    assert _derive_risk_level("Desconocida", "Raro") == RiskLevel.MEDIO


# ---------------------------------------------------------------------------
# _extract_json_array
# ---------------------------------------------------------------------------

def test_extract_json_array_plain() -> None:
    raw = '[{"a": 1}]'
    result = _extract_json_array(raw)
    assert result == [{"a": 1}]


def test_extract_json_array_with_markdown_fence() -> None:
    raw = "```json\n[{\"a\": 1}]\n```"
    result = _extract_json_array(raw)
    assert result == [{"a": 1}]


def test_extract_json_array_no_array_raises() -> None:
    with pytest.raises(ValueError, match="No JSON array"):
        _extract_json_array("No hay JSON aquí")


# ---------------------------------------------------------------------------
# RiskMatrixGenerator.generate
# ---------------------------------------------------------------------------

def test_generate_happy_path() -> None:
    client = _make_client(_SAMPLE_JSON)
    gen = RiskMatrixGenerator(client)
    consult = _make_consult_response()

    matrix = gen.generate(consult, query="Riesgo capital y DORA")

    assert matrix.trace_id == "test-trace-001"
    assert matrix.query == "Riesgo capital y DORA"
    assert len(matrix.items) == 3

    item0 = matrix.items[0]
    assert item0.regulation == "CRR Art. 92"
    assert item0.risk_level == RiskLevel.ALTO   # Alta × Alto
    assert item0.likelihood == "Alta"
    assert item0.impact == "Alto"

    item2 = matrix.items[2]
    assert item2.risk_level == RiskLevel.BAJO   # Baja × Bajo


def test_generate_llm_empty_returns_empty_matrix() -> None:
    client = _make_client("")
    gen = RiskMatrixGenerator(client)
    consult = _make_consult_response()

    matrix = gen.generate(consult, query="Cualquier consulta")
    assert matrix.items == []


def test_generate_llm_malformed_json_returns_empty_matrix() -> None:
    client = _make_client("esto no es json {{{")
    gen = RiskMatrixGenerator(client)
    consult = _make_consult_response()

    matrix = gen.generate(consult, query="Consulta")
    assert matrix.items == []


def test_generate_llm_raises_returns_empty_matrix() -> None:
    client = MagicMock()
    client.messages_create.side_effect = RuntimeError("API timeout")
    gen = RiskMatrixGenerator(client)
    consult = _make_consult_response()

    matrix = gen.generate(consult, query="Consulta")
    assert matrix.items == []


def test_generate_partial_items_skips_bad_entries() -> None:
    items = [
        {
            "regulation": "OK Art. 1",
            "obligation": "Obligación válida",
            "likelihood": "Media",
            "impact": "Alto",
            "mitigation": "Acción",
        },
        {
            # Bad likelihood value — Pydantic literal validation fails
            "regulation": "BAD Art. 2",
            "obligation": "Obligación inválida",
            "likelihood": "INVALID",
            "impact": "Alto",
            "mitigation": "Acción",
        },
    ]
    client = _make_client(json.dumps(items))
    gen = RiskMatrixGenerator(client)
    consult = _make_consult_response()

    matrix = gen.generate(consult, query="Test")
    # Bad entry is skipped; good one survives
    assert len(matrix.items) == 1
    assert matrix.items[0].regulation == "OK Art. 1"


# ---------------------------------------------------------------------------
# RiskMatrix model
# ---------------------------------------------------------------------------

def test_risk_matrix_generated_at_defaults_to_utc_now() -> None:
    before = datetime.now(UTC)
    matrix = RiskMatrix(trace_id="t1", query="q", items=[])
    after = datetime.now(UTC)
    assert before <= matrix.generated_at <= after


# ---------------------------------------------------------------------------
# build_risk_matrix_xlsx
# ---------------------------------------------------------------------------

def _make_matrix(n_items: int = 3) -> RiskMatrix:
    items = [
        RiskItem(
            regulation=f"Norma {i}",
            obligation=f"Obligación {i}",
            risk_level=RiskLevel.ALTO if i % 3 == 0 else (RiskLevel.MEDIO if i % 3 == 1 else RiskLevel.BAJO),
            likelihood="Alta" if i % 3 == 0 else ("Media" if i % 3 == 1 else "Baja"),
            impact="Alto" if i % 3 == 0 else ("Medio" if i % 3 == 1 else "Bajo"),
            mitigation=f"Mitigación {i}",
        )
        for i in range(n_items)
    ]
    return RiskMatrix(
        trace_id="xlsx-test-001",
        query="Consulta XLSX",
        items=items,
        generated_at=datetime(2026, 5, 18, 10, 0, 0, tzinfo=UTC),
    )


def test_build_risk_matrix_xlsx_returns_bytes() -> None:
    matrix = _make_matrix()
    result = build_risk_matrix_xlsx(matrix)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_build_risk_matrix_xlsx_is_valid_xlsx() -> None:
    """Check the magic bytes: XLSX is a ZIP — starts with PK\x03\x04."""
    matrix = _make_matrix()
    result = build_risk_matrix_xlsx(matrix)
    assert result[:4] == b"PK\x03\x04"


def test_build_risk_matrix_xlsx_empty_items() -> None:
    matrix = _make_matrix(n_items=0)
    result = build_risk_matrix_xlsx(matrix)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_build_risk_matrix_xlsx_sheet_names() -> None:
    """Verify both expected sheets exist in the workbook."""
    from io import BytesIO

    import openpyxl

    matrix = _make_matrix(3)
    result = build_risk_matrix_xlsx(matrix)
    wb = openpyxl.load_workbook(BytesIO(result))
    assert "Matriz de Riesgos" in wb.sheetnames
    assert "Resumen" in wb.sheetnames


def test_build_risk_matrix_xlsx_row_count() -> None:
    """Sheet 1 must have header + n_items rows."""
    from io import BytesIO

    import openpyxl

    n = 5
    matrix = _make_matrix(n)
    result = build_risk_matrix_xlsx(matrix)
    wb = openpyxl.load_workbook(BytesIO(result))
    ws = wb["Matriz de Riesgos"]
    # max_row includes header row
    assert ws.max_row == n + 1
