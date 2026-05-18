"""XLSX export for RiskMatrix — color-coded by risk level.

Uses openpyxl to produce a two-sheet workbook:
  - Sheet 1 "Matriz de Riesgos": one row per RiskItem, RAG colour coding.
  - Sheet 2 "Resumen": counts per level + generation metadata.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from lex_agents_agents.risk_matrix import RiskLevel, RiskMatrix

# ---------------------------------------------------------------------------
# Colour palette (matches Santander internal RAG standard)
# ---------------------------------------------------------------------------

_FILL_HEADER = PatternFill(fill_type="solid", fgColor="1A56DB")   # brand blue
_FILL_ALTO = PatternFill(fill_type="solid", fgColor="FEE2E2")     # red-100
_FILL_MEDIO = PatternFill(fill_type="solid", fgColor="FEF3C7")    # amber-100
_FILL_BAJO = PatternFill(fill_type="solid", fgColor="D1FAE5")     # green-100
_FILL_SUMMARY_ALTO = PatternFill(fill_type="solid", fgColor="FCA5A5")  # red-300
_FILL_SUMMARY_MEDIO = PatternFill(fill_type="solid", fgColor="FCD34D")  # amber-300
_FILL_SUMMARY_BAJO = PatternFill(fill_type="solid", fgColor="6EE7B7")   # green-300

_FONT_HEADER = Font(bold=True, color="FFFFFF", size=11)
_FONT_BOLD = Font(bold=True, size=10)
_FONT_NORMAL = Font(size=10)

_LEVEL_FILL = {
    RiskLevel.ALTO: _FILL_ALTO,
    RiskLevel.MEDIO: _FILL_MEDIO,
    RiskLevel.BAJO: _FILL_BAJO,
}

_COLUMNS = [
    ("Regulación", 35),
    ("Obligación", 55),
    ("Nivel de Riesgo", 18),
    ("Probabilidad", 16),
    ("Impacto", 14),
    ("Mitigación", 55),
]


def build_risk_matrix_xlsx(matrix: RiskMatrix) -> bytes:
    """Build an XLSX workbook from *matrix* and return the raw bytes."""
    wb = Workbook()

    # ── Sheet 1: matrix ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Matriz de Riesgos"

    # Header row
    for col_idx, (label, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = _FILL_HEADER
        cell.font = _FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30

    # Data rows
    for row_idx, item in enumerate(matrix.items, start=2):
        row_fill = _LEVEL_FILL.get(item.risk_level, _FILL_BAJO)
        values = [
            item.regulation,
            item.obligation,
            item.risk_level.value,
            item.likelihood,
            item.impact,
            item.mitigation,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.font = _FONT_NORMAL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = 40

    # Freeze header
    ws.freeze_panes = "A2"

    # Auto filter
    ws.auto_filter.ref = ws.dimensions

    # ── Sheet 2: resumen ─────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Resumen")

    counts = {
        RiskLevel.ALTO: sum(1 for i in matrix.items if i.risk_level == RiskLevel.ALTO),
        RiskLevel.MEDIO: sum(1 for i in matrix.items if i.risk_level == RiskLevel.MEDIO),
        RiskLevel.BAJO: sum(1 for i in matrix.items if i.risk_level == RiskLevel.BAJO),
    }

    summary_rows = [
        ("Consulta", matrix.query),
        ("Trace ID", matrix.trace_id),
        ("Generado", matrix.generated_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("Total items", len(matrix.items)),
        ("", ""),
        ("Nivel", "Cantidad"),
        ("Alto", counts[RiskLevel.ALTO]),
        ("Medio", counts[RiskLevel.MEDIO]),
        ("Bajo", counts[RiskLevel.BAJO]),
    ]

    summary_fills = {6: _FILL_HEADER, 7: _FILL_SUMMARY_ALTO, 8: _FILL_SUMMARY_MEDIO, 9: _FILL_SUMMARY_BAJO}
    summary_fonts = {6: _FONT_HEADER}

    for row_idx, (label, value) in enumerate(summary_rows, start=1):  # type: ignore[assignment]
        c_label = ws2.cell(row=row_idx, column=1, value=label)
        c_value = ws2.cell(row=row_idx, column=2, value=value)
        fill = summary_fills.get(row_idx)
        font = summary_fonts.get(row_idx, _FONT_NORMAL)
        if fill:
            c_label.fill = fill
            c_value.fill = fill
        if row_idx == 6:
            c_label.font = _FONT_HEADER
            c_value.font = _FONT_HEADER
        else:
            c_label.font = _FONT_BOLD
            c_value.font = font
        c_label.alignment = Alignment(horizontal="left")
        c_value.alignment = Alignment(horizontal="left", wrap_text=True)

    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 60

    # Write to bytes
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
