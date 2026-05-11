from __future__ import annotations
import hashlib, io
import docx  # python-docx
from lex_agents_documents.exceptions import ParseError
from lex_agents_documents.types import ParsedDocument

class DocxParser:
    def parse(self, raw_bytes: bytes, filename: str, doc_id: str) -> ParsedDocument:
        try:
            document = docx.Document(io.BytesIO(raw_bytes))
        except Exception as e:
            raise ParseError(f"corrupt: {e}") from e

        lines: list[str] = []
        for para in document.paragraphs:
            if para.style.name.startswith("Heading"):
                lines.append(f"\n## {para.text}")
            else:
                lines.append(para.text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                lines.append("| " + " | ".join(cells) + " |")

        text = "\n".join(lines)
        return ParsedDocument(
            doc_id=doc_id,
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            sha256=hashlib.sha256(raw_bytes).hexdigest(),
            text=text,
        )
