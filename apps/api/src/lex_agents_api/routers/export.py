"""Export and feedback endpoints for consultations."""

from __future__ import annotations

import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import structlog
from docx import Document
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from pydantic import BaseModel

from lex_agents_api.auth import CurrentUser, require_auth
from lex_agents_api.db import ConsultationStore
from lex_agents_api.routers.consult import get_store

logger: structlog.BoundLogger = structlog.get_logger(__name__)

router = APIRouter(prefix="/api/v1/consult", tags=["export"])

_DISCLAIMER = (
    "AVISO LEGAL: Este documento es un borrador generado asistidamente por IA "
    "(lex-agents v0.1.0). Requiere validación por un jurista cualificado antes de "
    "cualquier uso externo. No constituye asesoramiento jurídico profesional. "
    "Las citas deben verificarse en las fuentes originales (EUR-Lex, BOE)."
)


# ---------------------------------------------------------------------------
# DOCX export
# ---------------------------------------------------------------------------

def _build_docx(record_json: str, query: str) -> bytes:
    data: Any = json.loads(record_json)
    answer: str = str(data.get("answer", ""))
    citations: list[Any] = data.get("citations", [])
    metadata: dict[str, Any] = data.get("metadata", {})
    verification: dict[str, Any] | None = data.get("verification")

    doc = Document()

    # Cover / header
    title = doc.add_heading("Análisis jurídico asistido — lex-agents", level=0)
    title.runs[0].font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)  # brand blue

    doc.add_paragraph(f"Consulta: {query}")
    doc.add_paragraph(f"Fecha: {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}")
    doc.add_paragraph(f"Modelo: {metadata.get('model', 'desconocido')}")
    doc.add_paragraph(f"Prompt version: {metadata.get('prompt_version', '-')}")

    # Disclaimer banner
    p = doc.add_paragraph()
    run = p.add_run(_DISCLAIMER)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xB4, 0x5A, 0x00)  # amber

    doc.add_heading("Análisis", level=1)
    doc.add_paragraph(answer)

    # Verification status
    if verification:
        status = verification.get("status", "unknown")
        doc.add_heading("Estado de verificación", level=2)
        doc.add_paragraph(
            f"Estado: {status.upper()}  |  "
            f"Claims verificados: {verification.get('claims_passed', 0)}/{verification.get('claims_total', 0)}  |  "
            f"Referencias rotas: {len(verification.get('broken_refs', []))}"
        )

    # Citations
    if citations:
        doc.add_heading("Referencias normativas", level=1)
        for i, cit in enumerate(citations, 1):
            doc.add_heading(f"[REF:{cit.get('index', i)}] {cit.get('hierarchy_path', '')}", level=3)
            doc.add_paragraph(f"Fuente: {cit.get('source_id', '')}")
            fragment = cit.get("fragment_text", "")
            if fragment:
                p = doc.add_paragraph(style="Quote")
                p.add_run(fragment[:1000])

    # Technical metadata
    doc.add_heading("Metadatos técnicos", level=1)
    doc.add_paragraph(f"Latencia: {metadata.get('latency_ms', '-')} ms")
    doc.add_paragraph(f"Coste estimado: ${metadata.get('cost_estimate_usd', 0):.4f}")

    # Footer disclaimer on last paragraph
    doc.add_paragraph()
    doc.add_paragraph(_DISCLAIMER)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.get("/{trace_id}/export")
async def export_consultation(
    trace_id: str,
    _user: CurrentUser = Depends(require_auth),
    store: ConsultationStore = Depends(get_store),
) -> StreamingResponse:
    """Export a consultation as a .docx file."""
    record = await store.get(trace_id)
    if record is None:
        raise HTTPException(status_code=404, detail=f"Consultation {trace_id!r} not found")

    docx_bytes = _build_docx(record.response_json, record.query)
    filename = f"analisis_{trace_id[:8]}_{datetime.now(UTC).strftime('%Y%m%d')}.docx"

    logger.info("consultation_exported", trace_id=trace_id, format="docx")
    return StreamingResponse(
        BytesIO(docx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# XLSX comparative export
# ---------------------------------------------------------------------------

_FILL_HEADER = PatternFill(fill_type="solid", fgColor="1A56DB")
_FILL_INSUFFICIENT = PatternFill(fill_type="solid", fgColor="FFF3CD")
_FILL_PARTIAL = PatternFill(fill_type="solid", fgColor="FEF9C3")
_FILL_RISK_HIGH = PatternFill(fill_type="solid", fgColor="FEE2E2")
_FILL_RISK_MEDIUM = PatternFill(fill_type="solid", fgColor="FEF3C7")
_FILL_RISK_LOW = PatternFill(fill_type="solid", fgColor="D1FAE5")

_RISK_FILLS = {"high": _FILL_RISK_HIGH, "medium": _FILL_RISK_MEDIUM, "low": _FILL_RISK_LOW}


def _build_comparative_xlsx(comparative: dict[str, Any], query: str) -> bytes:
    wb = Workbook()

    # ── Sheet 1: Pivot table ────────────────────────────────────────────────
    ws_pivot = wb.active
    ws_pivot.title = "Análisis comparativo"

    jurisdictions: list[str] = comparative.get("jurisdictions_compared", [])
    dimensions: list[dict[str, Any]] = comparative.get("dimensions", [])

    # Title row
    ws_pivot.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + len(jurisdictions))
    title_cell = ws_pivot.cell(row=1, column=1, value=comparative.get("issue", query))
    title_cell.font = Font(bold=True, size=12, color="FFFFFF")
    title_cell.fill = _FILL_HEADER
    title_cell.alignment = Alignment(wrap_text=True)

    # Header row: Dimensión | J1 | J2 | …
    header_font = Font(bold=True, color="FFFFFF")
    ws_pivot.cell(row=2, column=1, value="Dimensión").font = header_font
    ws_pivot.cell(row=2, column=1).fill = _FILL_HEADER
    for col, j in enumerate(jurisdictions, start=2):
        cell = ws_pivot.cell(row=2, column=col, value=j)
        cell.font = header_font
        cell.fill = _FILL_HEADER
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for row_idx, dim in enumerate(dimensions, start=3):
        ws_pivot.cell(row=row_idx, column=1, value=dim.get("name", "")).font = Font(bold=True)
        by_j: dict[str, Any] = dim.get("by_jurisdiction", {})
        for col, j in enumerate(jurisdictions, start=2):
            entry: dict[str, Any] = by_j.get(j, {})
            coverage = entry.get("coverage", "full")
            text = entry.get("text") or "—"
            note = entry.get("note")
            refs = entry.get("refs", [])
            ref_str = " ".join(f"[REF:{r}]" for r in refs) if refs else ""
            cell_val = f"{text}\n{ref_str}".strip()
            if note:
                cell_val += f"\n⚠ {note}"
            cell = ws_pivot.cell(row=row_idx, column=col, value=cell_val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if coverage == "insufficient":
                cell.fill = _FILL_INSUFFICIENT
            elif coverage == "partial":
                cell.fill = _FILL_PARTIAL

    # Column widths
    ws_pivot.column_dimensions["A"].width = 28
    for col in range(2, 2 + len(jurisdictions)):
        col_letter = ws_pivot.cell(row=1, column=col).column_letter
        ws_pivot.column_dimensions[col_letter].width = 45

    # ── Sheet 2: Divergences ────────────────────────────────────────────────
    ws_div = wb.create_sheet("Divergencias")
    div_headers = ["Dimensión", "Descripción", "Jurisdicciones", "Severidad"]
    for col, h in enumerate(div_headers, start=1):
        cell = ws_div.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _FILL_HEADER

    for row_idx, div in enumerate(comparative.get("divergences", []), start=2):
        severity = div.get("severity", "medium")
        ws_div.cell(row=row_idx, column=1, value=div.get("dimension", ""))
        ws_div.cell(row=row_idx, column=2, value=div.get("description", "")).alignment = Alignment(wrap_text=True)
        ws_div.cell(row=row_idx, column=3, value=", ".join(div.get("jurisdictions_involved", [])))
        sev_cell = ws_div.cell(row=row_idx, column=4, value=severity.upper())
        sev_cell.fill = _RISK_FILLS.get(severity, PatternFill())
        sev_cell.alignment = Alignment(horizontal="center")

    for col_letter, width in [("A", 28), ("B", 55), ("C", 20), ("D", 12)]:
        ws_div.column_dimensions[col_letter].width = width

    # ── Sheet 3: Risk differential ──────────────────────────────────────────
    ws_risk = wb.create_sheet("Riesgo diferencial")
    ws_risk.cell(row=1, column=1, value="Jurisdicción").font = Font(bold=True, color="FFFFFF")
    ws_risk.cell(row=1, column=1).fill = _FILL_HEADER
    ws_risk.cell(row=1, column=2, value="Nivel de riesgo").font = Font(bold=True, color="FFFFFF")
    ws_risk.cell(row=1, column=2).fill = _FILL_HEADER

    risk_diff: dict[str, str] = comparative.get("risk_differential", {})
    for row_idx, j in enumerate(jurisdictions, start=2):
        level = risk_diff.get(j, "unknown")
        ws_risk.cell(row=row_idx, column=1, value=j)
        risk_cell = ws_risk.cell(row=row_idx, column=2, value=level.upper())
        risk_cell.fill = _RISK_FILLS.get(level, PatternFill())
        risk_cell.alignment = Alignment(horizontal="center")

    ws_risk.cell(row=len(jurisdictions) + 3, column=1, value=comparative.get("risk_rationale", ""))
    ws_risk.cell(row=len(jurisdictions) + 3, column=1).alignment = Alignment(wrap_text=True)
    ws_risk.column_dimensions["A"].width = 20
    ws_risk.column_dimensions["B"].width = 20
    ws_risk.merge_cells(
        start_row=len(jurisdictions) + 3,
        start_column=1,
        end_row=len(jurisdictions) + 3,
        end_column=2,
    )

    # ── Sheet 4: Citations ──────────────────────────────────────────────────
    ws_cit = wb.create_sheet("Citas")
    cit_headers = ["#", "Fuente", "Jerarquía", "Fragmento"]
    for col, h in enumerate(cit_headers, start=1):
        cell = ws_cit.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _FILL_HEADER

    for row_idx, cit in enumerate(comparative.get("citations", []), start=2):
        ws_cit.cell(row=row_idx, column=1, value=cit.get("index", row_idx - 1))
        ws_cit.cell(row=row_idx, column=2, value=cit.get("source_id", ""))
        ws_cit.cell(row=row_idx, column=3, value=cit.get("hierarchy_path", ""))
        frag_cell = ws_cit.cell(row=row_idx, column=4, value=cit.get("fragment_text", "")[:500])
        frag_cell.alignment = Alignment(wrap_text=True)

    for col_letter, width in [("A", 6), ("B", 25), ("C", 40), ("D", 70)]:
        ws_cit.column_dimensions[col_letter].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/{trace_id}/export/comparative")
async def export_comparative(
    trace_id: str,
    _user: CurrentUser = Depends(require_auth),
    store: ConsultationStore = Depends(get_store),
) -> StreamingResponse:
    """Export comparative analysis as XLSX (pivot table, divergences, risk, citations)."""
    record = await store.get(trace_id)
    if record is None:
        raise HTTPException(status_code=404, detail=f"Consultation {trace_id!r} not found")

    data: dict[str, Any] = json.loads(record.response_json)
    comparative = data.get("comparative_output")
    if comparative is None:
        raise HTTPException(
            status_code=422,
            detail="This consultation has no comparative analysis. Use output_type=analisis_comparativo.",
        )

    xlsx_bytes = _build_comparative_xlsx(comparative, record.query)
    filename = f"comparativo_{trace_id[:8]}_{datetime.now(UTC).strftime('%Y%m%d')}.xlsx"

    logger.info("comparative_exported", trace_id=trace_id, format="xlsx")
    return StreamingResponse(
        BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Feedback
# ---------------------------------------------------------------------------

class FeedbackRequest(BaseModel):
    issue_type: str  # error_factual | cita_incorrecta | fuera_de_alcance | otro
    description: str
    answer_excerpt: str = ""


@router.post("/{trace_id}/feedback", status_code=201)
async def submit_feedback(
    trace_id: str,
    body: FeedbackRequest,
    _user: CurrentUser = Depends(require_auth),
    store: ConsultationStore = Depends(get_store),
) -> dict[str, str]:
    """Save user feedback about a consultation to evals/feedback/."""
    record = await store.get(trace_id)
    if record is None:
        raise HTTPException(status_code=404, detail=f"Consultation {trace_id!r} not found")

    feedback_dir = Path("evals/feedback")
    feedback_dir.mkdir(parents=True, exist_ok=True)

    ts = datetime.now(UTC).strftime("%Y%m%dT%H%M%S")
    filename = feedback_dir / f"{ts}_{trace_id[:8]}.json"

    payload = {
        "trace_id": trace_id,
        "timestamp": datetime.now(UTC).isoformat(),
        "issue_type": body.issue_type,
        "description": body.description,
        "query": record.query,
        "answer_excerpt": body.answer_excerpt[:500],
    }
    filename.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    logger.info("feedback_saved", trace_id=trace_id, issue_type=body.issue_type, path=str(filename))
    return {"status": "saved", "path": str(filename)}
